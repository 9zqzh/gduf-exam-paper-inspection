#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""提取 docx / xlsx / pdf / txt 文本，供 Agent 深度核查试题、分析表、成绩单。

用法：
    python extract_text.py "材料/期末考试试题A卷.docx"
    python extract_text.py "材料/课程试卷及成绩分析表.docx" --max-chars 20000
    python extract_text.py "材料/试卷.pdf" --out "临时/试题.txt"

说明：
- 扫描版 PDF 通常无可提取文字，请直接用 Read 工具按页查看图像。
- 输出含字符数统计，便于核对分析表“每段约200字”等要求。
"""

import argparse
import json
import os
import sys

sys.stdout.reconfigure(encoding="utf-8")
sys.stderr.reconfigure(encoding="utf-8")


def from_docx(path):
    import docx
    d = docx.Document(path)
    parts = []
    for p in d.paragraphs:
        t = p.text.strip()
        if t:
            parts.append(t)
    for ti, table in enumerate(d.tables, 1):
        parts.append(f"\n[表 {ti}]")
        for row in table.rows:
            cells = []
            prev = None
            for c in row.cells:
                txt = c.text.strip().replace("\n", " ")
                if txt != prev:
                    cells.append(txt)
                prev = txt
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def from_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"\n[工作表：{ws.title}]")
        for row in ws.iter_rows(values_only=True):
            if not any(v not in (None, "") for v in row):
                continue
            parts.append(" | ".join("" if v is None else str(v).replace("\n", " ") for v in row))
    wb.close()
    return "\n".join(parts)


def from_pdf(path):
    reader = None
    try:
        import pypdf
        reader = pypdf.PdfReader(path)
    except ImportError:
        try:
            import PyPDF2
            reader = PyPDF2.PdfReader(path)
        except ImportError:
            return ("[未能提取 PDF 文本：未安装 pypdf。请执行 python -m pip install pypdf，"
                    "或直接用 Read 工具按页查看；扫描件请用视觉方式核对。]")
    parts = []
    for i, page in enumerate(reader.pages, 1):
        try:
            t = page.extract_text() or ""
        except Exception as exc:  # noqa: BLE001
            t = f"[第{i}页提取失败：{exc}]"
        parts.append(f"\n[第 {i} 页]\n{t.strip()}")
    text = "\n".join(parts)
    if len(text.strip()) < 50:
        text += "\n[提示：该 PDF 可能是扫描件，几乎无可提取文字，请改用 Read 工具按页查看图像。]"
    return text


def extract(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".docx", ".doc"):
        if ext == ".doc":
            return "[.doc 旧格式无法直接解析，请另存为 .docx 后重试，或用 Word 打开复制内容。]"
        return from_docx(path)
    if ext in (".xlsx", ".xlsm"):
        return from_xlsx(path)
    if ext == ".pdf":
        return from_pdf(path)
    if ext in (".txt", ".md", ".csv"):
        for enc in ("utf-8", "gbk", "utf-16"):
            try:
                with open(path, encoding=enc) as fh:
                    return fh.read()
            except UnicodeDecodeError:
                continue
        return "[文本编码无法识别，请转存为 UTF-8。]"
    return f"[不支持的文件类型：{ext}]"


def main():
    ap = argparse.ArgumentParser(description="提取材料文本")
    ap.add_argument("path", help="文件路径")
    ap.add_argument("--max-chars", type=int, default=60000, help="最多输出字符数，默认 60000")
    ap.add_argument("--out", help="同时写入文本文件")
    ap.add_argument("--json", action="store_true", help="以 JSON 输出（含字符数）")
    args = ap.parse_args()

    if not os.path.isfile(args.path):
        raise SystemExit(f"文件不存在：{args.path}")
    text = extract(args.path)
    full_len = len(text)
    shown = text[: args.max_chars]
    if full_len > args.max_chars:
        shown += f"\n\n[已截断：原文 {full_len} 字符，仅显示前 {args.max_chars} 字符]"

    if args.out:
        os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
        with open(args.out, "w", encoding="utf-8") as fh:
            fh.write(text)

    if args.json:
        print(json.dumps({
            "path": os.path.abspath(args.path),
            "chars": full_len,
            "text": shown,
        }, ensure_ascii=False, indent=2))
    else:
        print(f"[文件] {args.path}")
        print(f"[字符数] {full_len}")
        print("-" * 40)
        print(shown)


if __name__ == "__main__":
    main()
