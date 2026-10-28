#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""按学院《教师试卷检查结果报告》模板排版生成 Word 文档。

用法：
    python build_report_docx.py --payload 输出/report_payload.json --out "输出/张三_试卷检查结果报告.docx"

payload 结构（UTF-8 JSON）：
{
  "teacher": "张三",
  "department": "计算机与科学系",
  "course_name": "Java程序设计",
  "course_type": "必修课",              // 必修课/选修课/课程论文/实验（实践）课
  "term": "2025-2026-2学期",
  "assessment": "闭卷笔试；过程性考核占比30%",
  "classes": "25清远计算机科学与技术-1～3班，共167人",
  "paper_sets": "A/B卷各1套",
  "subtitle": "（2025—2026学年试卷专项检查）",   // 可选
  "sections": [
    {"heading": "一、课程与试卷基本情况", "paras": ["……", "……"]},
    {"heading": "二、命题规范与“两性一度”自查情况", "paras": ["（一）……", "……"]},
    {"heading": "三、成绩统计分析", "paras": ["……"],
     "table": {"headers": ["项目","平均分","及格率","优秀率","不及格率","标准差"],
               "rows": [["最终成绩","78.4","92%","18%","8%","9.6"]]}}
  ],
  "sign_date": "2026年11月3日"          // 可选，默认留空待手写
}

脚本会输出“Word 式字数”统计；不足 1000 字时以非零退出码提示。
"""

import argparse
import json
import os
import re
import sys

try:
    import docx
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
except ImportError:
    sys.stderr.write("ERROR: 缺少 python-docx，请执行: python -m pip install python-docx\n")
    sys.exit(2)

sys.stdout.reconfigure(encoding="utf-8")
sys.stderr.reconfigure(encoding="utf-8")

COURSE_TYPES = ["必修课", "选修课", "课程论文", "实验（实践）课"]


def set_run_font(run, name_cn="宋体", name_ascii="Times New Roman", size=12, bold=False):
    run.font.name = name_ascii
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name_cn)
    rfonts.set(qn("w:ascii"), name_ascii)
    rfonts.set(qn("w:hAnsi"), name_ascii)


def add_para(doc, text, size=12, bold=False, align=None, indent=True,
             before=0, after=0, line=1.5, cn="宋体", ascii_font="Times New Roman"):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if align is not None:
        p.alignment = align
    if indent:
        pf.first_line_indent = Pt(size * 2)
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = line
    run = p.add_run(text)
    set_run_font(run, cn, ascii_font, size, bold)
    return p


def shade_cell(cell, color="F2F2F2"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def style_table(table, widths=None, font_size=10.5, header_bold=True):
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths and idx < len(widths):
                cell.width = widths[idx]
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.2
                for r in p.runs:
                    set_run_font(r, "宋体", "Times New Roman", font_size, header_bold and row._index == 0)
    if widths:
        for idx, w in enumerate(widths):
            table.columns[idx].width = w


def count_words(text):
    cjk = len(re.findall(r"[\u4e00-\u9fff]", text))
    latin = len(re.findall(r"[A-Za-z0-9]+", text))
    return cjk, latin, cjk + latin


def build(payload, out_path):
    doc = docx.Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(3.175)
    sec.right_margin = Cm(3.175)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    # 标题
    add_para(doc, "教师试卷检查结果报告", size=16, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=6, cn="黑体")
    if payload.get("subtitle"):
        add_para(doc, payload["subtitle"], size=12, align=WD_ALIGN_PARAGRAPH.CENTER,
                 indent=False, after=12)

    # 基本信息表
    table = doc.add_table(rows=8, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    labels = ["任课教师", "所属教研室", "课程名称", "课程类型", "学年学期", "考核方式", "班级及人数", "试卷套数"]
    values = [
        payload.get("teacher", ""),
        payload.get("department", ""),
        payload.get("course_name", ""),
        None,  # 课程类型单独渲染复选框
        payload.get("term", ""),
        payload.get("assessment", ""),
        payload.get("classes", ""),
        payload.get("paper_sets", ""),
    ]
    widths = [Cm(3.78), Cm(11.25)]
    for i, (label, value) in enumerate(zip(labels, values)):
        c0, c1 = table.rows[i].cells
        c0.text = ""
        c0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c0.paragraphs[0].add_run(label)
        set_run_font(r, "宋体", "Times New Roman", 10.5, True)
        c1.text = ""
        p1 = c1.paragraphs[0]
        if label == "课程类型":
            selected = payload.get("course_type", "")
            text = "　".join(("■" if t == selected else "□") + t for t in COURSE_TYPES)
        else:
            text = str(value or "")
        r = p1.add_run(text)
        set_run_font(r, "宋体", "Times New Roman", 10.5, False)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = widths[idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.2
    for idx, w in enumerate(widths):
        table.columns[idx].width = w

    # 正文
    all_text = []
    for section in payload.get("sections", []):
        add_para(doc, section.get("heading", ""), size=14, bold=True,
                 indent=False, before=14, after=6, cn="黑体")
        for para in section.get("paras", []):
            if not str(para).strip():
                continue
            all_text.append(str(para))
            add_para(doc, str(para), size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        tbl = section.get("table")
        if tbl and tbl.get("headers"):
            headers = tbl["headers"]
            t = doc.add_table(rows=1, cols=len(headers))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False
            ncol = len(headers)
            # 第一列放“统计口径/项目”，适当加宽，其余列均分正文宽度
            total_w = 14.6
            first_w = 3.4 if ncol >= 5 else total_w / max(ncol, 1)
            rest_w = (total_w - first_w) / (ncol - 1) if ncol > 1 else total_w
            col_widths = [Cm(first_w)] + [Cm(rest_w)] * (ncol - 1)
            for j, w in enumerate(col_widths):
                t.columns[j].width = w
            for j, h in enumerate(headers):
                c = t.rows[0].cells[j]
                c.text = ""
                p = c.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(str(h))
                set_run_font(r, "宋体", "Times New Roman", 10.5, True)
                shade_cell(c)
            for row in tbl.get("rows", []):
                cells = t.add_row().cells
                for j, v in enumerate(row):
                    if j >= len(cells):
                        break
                    cells[j].text = ""
                    p = cells[j].paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r = p.add_run(str(v))
                    set_run_font(r, "宋体", "Times New Roman", 10.5, False)
            for row in t.rows:
                for j, cell in enumerate(row.cells):
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    if j < len(col_widths):
                        cell.width = col_widths[j]
                    for p in cell.paragraphs:
                        p.paragraph_format.space_before = Pt(2)
                        p.paragraph_format.space_after = Pt(2)
            add_para(doc, "", indent=False, after=0)

    # 签名
    add_para(doc, "", indent=False, before=12, after=0)
    add_para(doc, "教师签名：＿＿＿＿＿＿＿＿", size=12, indent=False,
             align=WD_ALIGN_PARAGRAPH.RIGHT, before=12)
    add_para(doc, f"日期：{payload.get('sign_date') or '2026年　　月　　日'}", size=12,
             indent=False, align=WD_ALIGN_PARAGRAPH.RIGHT)

    out_dir = os.path.dirname(os.path.abspath(out_path))
    os.makedirs(out_dir, exist_ok=True)
    doc.save(out_path)

    cjk, latin, total = count_words("\n".join(all_text))
    return {"out": os.path.abspath(out_path), "cjk": cjk, "words": latin, "total": total}


def main():
    ap = argparse.ArgumentParser(description="生成试卷检查结果报告 docx")
    ap.add_argument("--payload", required=True, help="JSON 数据文件")
    ap.add_argument("--out", required=True, help="输出 docx 路径")
    ap.add_argument("--min-words", type=int, default=1000, help="最少字数，默认1000")
    args = ap.parse_args()

    with open(args.payload, encoding="utf-8") as fh:
        payload = json.load(fh)
    res = build(payload, args.out)
    print(f"[已生成] {res['out']}")
    print(f"[字数] 汉字 {res['cjk']} + 英文/数字词 {res['words']} = Word 式字数约 {res['total']}")

    joined = json.dumps(payload, ensure_ascii=False)
    if "（填写" in joined or "TODO" in joined or "……" in joined:
        print("[提醒] 报告中可能仍含模板提示语或省略号，请检查后再交付。", file=sys.stderr)
    if res["total"] < args.min_words:
        print(f"[警告] 字数不足 {args.min_words} 字，请补充实质内容（不要用空话凑字数）。", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
